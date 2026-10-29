"""Auditoria somente leitura dos integrantes atuais da campanha."""

import csv
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from itertools import groupby
from pathlib import Path

from requests import RequestException

from analytics import cpf_query_priority, get_fields
from farmalider_client import (
    FarmaliderAuthError,
    FarmaliderError,
    FarmaliderUnavailableError,
)
from farmalider_customers import prepare_customers
from recompra import (
    RecompraReadCache,
    atomic_json,
    iter_prioritized_contacts,
    now_local,
    propose,
    resolve_latest,
)
from task_progress import TaskProgress


def audit_group(weni, farm, structure, path, *, now=None):
    now = now or now_local()
    contacts = list(
        iter_prioritized_contacts(weni.iter_group_contacts(structure["group_uuid"]))
    )
    prepare_customers(farm, len(contacts))
    reads = RecompraReadCache(farm)

    def check(contact):
        fields = contact.get("fields") or {}
        row = {
            "contact_uuid": contact["uuid"],
            "nome": contact.get("name", ""),
            "cpf": fields.get("cpf", ""),
            "pedido_weni": fields.get("pedido_id", ""),
            "data_weni": fields.get("pedido_data", ""),
            "status_weni": fields.get("pedido_status", ""),
            "pagamento_weni": fields.get("pedido_forma_pagamento", ""),
        }
        if "fields" not in contact:
            row.update(
                classificacao="CONSULTA_PENDENTE", motivo="CONTATO_DUPLICADO_DIVERGENTE"
            )
            return row
        try:
            order, items = resolve_latest(contact, reads, now.isoformat())
            plan = propose(contact, order, items, structure, now=now)
            order = order or {}
            reason = plan["reason"]
            pending = str(order.get("status_confirmacao", "")).startswith("ERRO_")
            row.update(
                pedido_api=order.get("pedido_id", ""),
                data_api=order.get("data_hora", ""),
                status_api=order.get("farmalider_status_pedido", ""),
                pagamento_api=order.get("farmalider_forma_pagamento", ""),
                motivo=reason,
                acao_proposta=plan["action"],
                classificacao="CONSULTA_PENDENTE"
                if pending
                else "DENTRO_DA_REGRA"
                if reason.startswith("JANELA_0_A_")
                else "REVISAR",
                campos_a_atualizar=", ".join(plan["fields"]),
            )
        except (FarmaliderAuthError, FarmaliderUnavailableError):
            raise
        except (FarmaliderError, RequestException) as exc:
            row.update(classificacao="CONSULTA_PENDENTE", motivo=type(exc).__name__)
        return row

    rows = []
    with (
        ThreadPoolExecutor(max_workers=4) as pool,
        TaskProgress(
            "Auditoria do grupo Recompra", len(contacts), batch_size=100, interval=15
        ) as progress,
    ):
        for _, phase in groupby(
            contacts, key=lambda c: cpf_query_priority(get_fields(c).get("cpf"))
        ):
            for row in pool.map(check, phase):
                rows.append(row)
                progress.advance()
    summary = {
        "somente_leitura": True,
        "segment": structure.get("segment", "FARMA"),
        "capturado_em": now.isoformat(),
        "contatos": len(rows),
        "classificacoes": dict(Counter(r["classificacao"] for r in rows)),
        "motivos": dict(Counter(r["motivo"] for r in rows)),
    }
    atomic_json(path, {"summary": summary, "contacts": rows})
    columns = [
        "nome",
        "cpf",
        "contact_uuid",
        "classificacao",
        "motivo",
        "pedido_weni",
        "data_weni",
        "status_weni",
        "pagamento_weni",
        "pedido_api",
        "data_api",
        "status_api",
        "pagamento_api",
        "acao_proposta",
        "campos_a_atualizar",
    ]
    with (
        Path(path)
        .with_suffix(".csv")
        .open("w", encoding="utf-8-sig", newline="") as file
    ):
        writer = csv.DictWriter(
            file, fieldnames=columns, extrasaction="ignore", delimiter=";"
        )
        writer.writeheader()
        writer.writerows(
            {
                key: (
                    "'" + value
                    if isinstance(value, str)
                    and value.lstrip().startswith(("=", "+", "-", "@"))
                    else value
                )
                for key, value in row.items()
            }
            for row in rows
        )
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    if sheet is None:
        raise RuntimeError("Não foi possível criar o relatório XLSX.")
    sheet.title = "Auditoria Recompra"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(key, "") for key in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    from openpyxl.utils import get_column_letter

    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 25
    book.save(Path(path).with_suffix(".xlsx"))
    return summary
