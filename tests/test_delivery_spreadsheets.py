"""Fixtures sintéticas, sem dados pessoais nem chamadas externas."""

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock

import openpyxl

from delivery_spreadsheets import (
    CONTACTS_FILENAME,
    CONVERSATIONS_FILENAME,
    DeliverySpreadsheetError,
    read_delivery_spreadsheets,
    validate_current_contacts,
)
from delivery_tickets import (
    apply_delivery_review,
    select_tickets,
    write_delivery_review,
)
from tests import test_delivery_tickets as fixtures
from tests.test_delivery_tickets import NOW, contact

CHAT = "30000000-0000-0000-0000-000000000001"
OTHER = "10000000-0000-0000-0000-000000000002"
URN = "whatsapp:5591987654321"
HEADERS = [
    "Contact UUID",
    "Name",
    "URN:Whatsapp",
    "URN:Whatsapp",
    "Field:CPF",
    "Field:Email",
    "Field:nome",
    "Field:Pedido Produtos",
]
ROOM_HEADERS = [
    "contact__uuid",
    "contact__name",
    "urn",
    "created_on",
    "ended_at",
    "is_active",
    "tags_list",
    "user__email",
    "queue__sector__name",
    "custom_fields",
]


class DeliverySpreadsheetsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.contact_rows = [
            [
                contact()["uuid"],
                "Cliente Árvore",
                "5591987654321",
                "",
                "0",
                "arvore@example.test",
                "Cliente Árvore",
                "Produto",
            ]
        ]
        self.room_rows = [
            [
                CHAT,
                "Cliente Arvore",
                URN,
                "2026-10-08 09:00:00",
                "2026-10-08 12:00:00",
                False,
                "['Pedido Finalizado']",
                "agent@example.test",
                "Delivery FarmaLíder",
                "{'pedido_id': '201'}",
            ]
        ]

    def exports(self):
        for name, sheet, headers, rows in [
            (CONTACTS_FILENAME, "Contact", HEADERS, self.contact_rows),
            (CONVERSATIONS_FILENAME, "rooms", ROOM_HEADERS, self.room_rows),
        ]:
            book = openpyxl.Workbook()
            page = book.active
            page.title = sheet
            page.append(headers)
            for row in rows:
                page.append(row)
            book.save(self.directory / name)

    def read(self):
        self.exports()
        return read_delivery_spreadsheets(
            self.directory, cutoff=datetime.fromisoformat(NOW)
        )

    def test_internal_uuid_maps_to_unique_weni_urn_and_full_name(self):
        tickets, contacts, report = self.read()
        self.assertEqual(tickets[0]["contact_uuid"], contact()["uuid"])
        self.assertEqual(tickets[0]["chat_contact_uuid"], CHAT)
        self.assertEqual(contacts[0]["fields"]["cpf"], "0")
        self.assertEqual(contacts[0]["fields"]["email"], "arvore@example.test")
        self.assertEqual(contacts[0]["fields"]["pedido_produtos"], "Produto")
        self.assertEqual(report["summary"]["clientes_vinculados"], 1)
        self.assertEqual(tickets[0]["custom_fields"], {})
        self.assertEqual(len(report["sources"][0]["sha256"]), 64)

    def test_deleted_order_tag_imports_only_exact_legacy_alias(self):
        for tag, expected in [
            ("Pedido is deleted", 1),
            (" PEDIDO  IS   DELETED ", 1),
            ("Pedido_is_deleted_2026-09-29 11:53:51.275814+00:00", 1),
            ("Pedido_is_deleted_2026-09-29 13:04:10.754801+00:00", 1),
            ("Pedido_is_deleted_2026-02-30T13:04:10Z", 0),
            ("Pedido_is_deleted_2026-09-29 13:04:10", 0),
            ("Orcamento_is_deleted_2026-09-29 11:53:51.275814+00:00", 0),
            ("Pedido is deleted extra", 0),
            ("Pedido cancelado", 0),
        ]:
            with self.subTest(tag=tag):
                self.room_rows[0][6] = repr([tag])
                self.assertEqual(len(self.read()[0]), expected)

    def test_repeated_urn_columns_are_preserved_and_conflicting_contacts_block(self):
        self.contact_rows[0][2] = "5591111111111"
        self.contact_rows[0][3] = "5591987654321"
        self.assertEqual(len(self.read()[0]), 1)
        self.contact_rows.append(
            [
                OTHER,
                "Cliente Árvore",
                "5591987654321",
                "",
                "0",
                "arvore@example.test",
                "",
                "",
            ]
        )
        tickets, _, report = self.read()
        self.assertEqual(tickets, [])
        self.assertEqual(report["summary"]["URN_AMBIGUA"], 1)

    def test_name_mismatch_and_non_native_field_urn_cannot_map(self):
        self.room_rows[0][1] = "Outra Pessoa"
        self.assertFalse(self.read()[0])
        self.room_rows[0][1] = "Cliente Árvore"
        self.room_rows[0][2] = "whatsapp:5591222222222"
        self.assertFalse(self.read()[0])

    def test_closed_tagged_human_delivery_only_and_time_cutoff(self):
        for index, value in [
            (5, True),
            (6, "['Pedidos']"),
            (7, ""),
            (8, "Outro setor"),
            (4, "2026-10-09 12:00:00"),
            (3, "2026-10-09 12:00:00"),
        ]:
            with self.subTest(index=index):
                old = self.room_rows[0][index]
                self.room_rows[0][index] = value
                self.assertFalse(self.read()[0])
                self.room_rows[0][index] = old

    def test_duplicates_and_api_export_timezone_have_same_stable_ticket_key(self):
        self.room_rows.append(list(self.room_rows[0]))
        tickets, _, report = self.read()
        self.assertEqual(len(tickets), 1)
        self.assertEqual(report["summary"]["ATENDIMENTO_DUPLICADO"], 1)
        live = {
            "uuid": OTHER,
            "contact": {"uuid": CHAT, "external_id": contact()["uuid"]},
            "tags": [{"name": "Pedido Finalizado"}],
            "user": {"name": "synthetic"},
            "created_on": "2026-10-08T12:00:00Z",
            "ended_at": "2026-10-08T15:00:00Z",
            "is_active": False,
        }
        selected, _ = select_tickets([live], cutoff=datetime.fromisoformat(NOW))
        self.assertEqual(selected[0]["ticket_uuid"], tickets[0]["ticket_uuid"])
        self.assertEqual(tickets[0]["aberto_em"], "2026-10-08T09:00:00-03:00")

    def test_duplicate_uuid_different_data_and_formulas_fail_closed(self):
        self.contact_rows.append(
            [contact()["uuid"], "Outra Pessoa", "", "", "0", "", "", ""]
        )
        with self.assertRaises(DeliverySpreadsheetError):
            self.read()
        self.contact_rows.pop()
        self.contact_rows[0][5] = "=1+1"
        with self.assertRaises(DeliverySpreadsheetError):
            self.read()

    def test_wrong_sheet_and_missing_file_fail_without_network(self):
        with self.assertRaises(DeliverySpreadsheetError):
            read_delivery_spreadsheets(
                self.directory, cutoff=datetime.fromisoformat(NOW)
            )
        self.exports()
        book = openpyxl.load_workbook(self.directory / CONTACTS_FILENAME)
        book.active.title = "Wrong"
        book.save(self.directory / CONTACTS_FILENAME)
        book.close()
        with self.assertRaises(DeliverySpreadsheetError):
            read_delivery_spreadsheets(
                self.directory, cutoff=datetime.fromisoformat(NOW)
            )

    def test_numbered_contact_tabs_from_real_export_are_all_read(self):
        self.exports()
        path = self.directory / CONTACTS_FILENAME
        book = openpyxl.load_workbook(path)
        book.active.title = "Contact 1"
        page = book.create_sheet("Contact 2")
        page.append(HEADERS)
        page.append(
            [
                OTHER,
                "Outro Cliente",
                "5591111111111",
                "",
                "52998224725",
                "other@example.test",
                "",
                "",
            ]
        )
        book.save(path)
        book.close()
        tickets, _, report = read_delivery_spreadsheets(
            self.directory, cutoff=datetime.fromisoformat(NOW)
        )
        self.assertEqual(len(tickets), 1)
        self.assertEqual(report["summary"]["contatos_export"], 2)

    def test_identical_api_pages_are_deduplicated(self):
        tickets, _, _ = self.read()
        current = contact(urns=[URN])
        refresh = Mock()
        accepted, contacts, _ = validate_current_contacts(
            tickets, [current, dict(current)], refresh_contact=refresh
        )
        self.assertEqual(len(accepted), 1)
        self.assertEqual(len(contacts), 1)
        refresh.assert_not_called()

    def test_changed_duplicate_is_refetched_and_old_urn_is_rejected(self):
        tickets, _, _ = self.read()
        current = contact(urns=[URN])
        changed = contact(urns=["whatsapp:5591222222222"])
        refresh = Mock(return_value=changed)
        accepted, _, rejected = validate_current_contacts(
            tickets, [current, changed], refresh_contact=refresh
        )
        refresh.assert_called_once_with(current["uuid"])
        self.assertFalse(accepted)
        self.assertEqual(rejected["URN_ALTERADA_OU_AMBIGUA_NA_WENI_ATUAL"], 1)

    def test_conflicting_pages_without_refetch_block(self):
        tickets, _, _ = self.read()
        current = contact(urns=[URN])
        changed = {**current, "name": "Outro Cliente"}
        with self.assertRaises(DeliverySpreadsheetError):
            validate_current_contacts(tickets, [current, changed])
        with self.assertRaises(DeliverySpreadsheetError):
            validate_current_contacts(
                tickets,
                [current, changed],
                refresh_contact=Mock(return_value={**current, "uuid": OTHER}),
            )

    def test_current_contact_uuid_native_urn_and_name_must_still_match(self):
        tickets, _, _ = self.read()
        current = contact(urns=[URN])
        selected, contacts, rejected = validate_current_contacts(tickets, [current])
        self.assertEqual(len(selected), 1)
        self.assertEqual(len(contacts), 1)
        self.assertEqual(rejected, {})
        changed = contact(urns=["whatsapp:5591222222222"])
        self.assertFalse(validate_current_contacts(tickets, [changed])[0])
        duplicate = contact(uuid=OTHER, urns=[URN])
        self.assertFalse(validate_current_contacts(tickets, [current, duplicate])[0])
        current["name"] = "Outra Pessoa"
        current["fields"]["nome"] = "Outra Pessoa"
        self.assertFalse(validate_current_contacts(tickets, [current])[0])
        self.assertFalse(validate_current_contacts(tickets, [])[0])

    def test_targets_are_frozen_weni_has_no_sheet_writes_analytics_has_no_weni_writes(
        self,
    ):
        fixture = fixtures.DeliveryTicketsTests()
        for target in ["weni", "analytics"]:
            with self.subTest(target=target):
                plans, data = fixture.reconcile()
                data["target"] = target
                path = self.directory / "review.json"
                write_delivery_review(plans, data, path)
                weni = Mock()
                weni.get_contact.return_value = contact()
                from weni_update import REQUIRED_ORDER_FIELDS

                weni.list_fields.return_value = [
                    {"key": key} for key in REQUIRED_ORDER_FIELDS
                ]
                store = Mock()
                for method in (
                    store.upsert_contacts_detailed,
                    store.upsert_orders_detailed,
                    store.upsert_order_items_detailed,
                ):
                    method.return_value.final_records = []
                store.upsert_pharmacy_orders.return_value = (
                    store.upsert_orders_detailed.return_value,
                    store.upsert_order_items_detailed.return_value,
                )
                store.reconcile_pharmacy_scope.return_value.orders = []
                store.reconcile_pharmacy_scope.return_value.items = []
                store.reconcile_pharmacy_scope.return_value.contacts = []
                factory = Mock(return_value=store)
                result = apply_delivery_review(
                    path, weni, store_factory=factory, target=target
                )
                self.assertEqual(result["contacts_accepted"], 1)
                if target == "weni":
                    weni.update_contact_fields.assert_called_once()
                    factory.assert_not_called()
                else:
                    weni.update_contact_fields.assert_not_called()
                    store.upsert_delivery_tickets.assert_called_once()
                weni.reset_mock()
                factory.reset_mock()
                with self.assertRaises(ValueError):
                    apply_delivery_review(
                        path, weni, store_factory=factory, target="both"
                    )
                weni.list_fields.assert_not_called()
                factory.assert_not_called()


if __name__ == "__main__":
    unittest.main()
